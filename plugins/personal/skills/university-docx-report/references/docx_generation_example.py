"""Minimal pattern for a Markdown-backed university DOCX report.

Use a copy of a style-only template as the starting point. Populate content
from the report's Markdown source; this file demonstrates style handling only.
Check the template numbering before export: when heading or caption styles
generate a number, pass only the title or caption text to add_paragraph.
Do not type the same number into the final document a second time.
"""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm


def require_style(document, name):
    if name not in {style.name for style in document.styles}:
        raise KeyError(f"Required template style is missing: {name}")
    return name


def clear_body(document):
    body = document._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def ensure_table_text_style(document, main_style_name):
    name = "P: Текст таблицы отчёта"
    if name not in {style.name for style in document.styles}:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles[main_style_name]
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return name


def ensure_reference_style(document, main_style_name):
    name = "P: Ссылка на таблицу или рисунок"
    if name not in {style.name for style in document.styles}:
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles[main_style_name]
        style.paragraph_format.first_line_indent = Cm(0)
    return name


def set_cell_text(cell, text, paragraph_style):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.style = paragraph_style
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.add_run(str(text))


def add_table(document, reference, caption, headers, rows, reference_style, table_text_style):
    document.add_paragraph(reference, style=reference_style)
    document.add_paragraph(caption, style=reference_style)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, table_text_style)
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            set_cell_text(cell, text, table_text_style)
    return table


def add_figure_placeholder(document, placeholder_text, illustration_style, caption=None, caption_style=None):
    document.add_paragraph(placeholder_text, style=illustration_style)
    if caption is not None:
        # For an auto-numbered caption style, strip "Рисунок 1.1 — " from the
        # approved Markdown caption before passing its remaining text here.
        document.add_paragraph(caption, style=caption_style)


def build(template_path, output_path):
    raise RuntimeError(
        "This is a style-only example with hardcoded text, not an approved report generator. "
        "Adapt it to read the canonical Markdown and call the critic gate before construction and saving."
    )
    document = Document(template_path)
    clear_body(document)
    main = require_style(document, "P: Мой основной текст")
    heading_1 = require_style(document, "P: Заголовок 1 уровня")
    illustration = require_style(document, "Иллюстрация")
    table_text = ensure_table_text_style(document, main)
    reference = ensure_reference_style(document, main)

    # The heading styles in "Просто стили" number paragraphs automatically.
    # A Markdown heading "1. АНАЛИЗ..." becomes text "АНАЛИЗ..." here.
    document.add_paragraph("АНАЛИЗ ПРЕДМЕТНОЙ ОБЛАСТИ", style=heading_1)
    document.add_paragraph("Тема", style="P: Заголовок 2 уровня")
    document.add_paragraph("Текст раздела.", style=main)
    add_table(
        document,
        "Исходные данные приведены в таблице 1.1.",
        "Таблица 1.1 — Исходные данные",
        ["Параметр", "Значение"],
        [("Параметр 1", "Значение 1")],
        reference,
        table_text,
    )
    add_figure_placeholder(
        document,
        "ВСТАВЬ СЮДА ИЗОБРАЖЕНИЕ ПРО МОДЕЛЬ СИСТЕМЫ",
        illustration,
    )
    document.save(output_path)


if __name__ == "__main__":
    build(Path("Просто стили.docx"), Path("Отчёт.docx"))
